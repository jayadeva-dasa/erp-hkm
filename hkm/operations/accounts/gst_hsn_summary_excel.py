import frappe
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import re
from datetime import datetime


def execute():
	"""Generate GST HSN Summary Excel Report"""
	
	print("Starting GST HSN Summary Excel generation...")
	
	# Fetch data
	data = get_gst_hsn_data()
	
	if not data:
		print("No data found matching the criteria")
		return
	
	# Create Excel file
	wb = openpyxl.Workbook()
	ws = wb.active
	ws.title = "GST HSN Summary"
	
	# Define headers
	headers = [
		"HSN Code",
		"UOM",
		"Quantity",
		"Taxable Amount",
		"Item Tax Template",
		"CGST",
		"SGST",
		"IGST",
		"Total Amount"
	]
	
	# Write headers
	for col_num, header in enumerate(headers, 1):
		cell = ws.cell(row=1, column=col_num)
		cell.value = header
		cell.font = Font(bold=True, color="FFFFFF")
		cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
		cell.alignment = Alignment(horizontal="center", vertical="center")
		cell.border = Border(
			left=Side(style='thin'),
			right=Side(style='thin'),
			top=Side(style='thin'),
			bottom=Side(style='thin')
		)
	
	# Write data rows
	row_num = 2
	for row_data in data:
		ws.cell(row=row_num, column=1, value=row_data.get("gst_hsn_code"))
		ws.cell(row=row_num, column=2, value=row_data.get("uom"))
		ws.cell(row=row_num, column=3, value=row_data.get("quantity"))
		ws.cell(row=row_num, column=4, value=row_data.get("taxable_amount"))
		ws.cell(row=row_num, column=5, value=row_data.get("item_tax_template"))
		ws.cell(row=row_num, column=6, value=row_data.get("cgst"))
		ws.cell(row=row_num, column=7, value=row_data.get("sgst"))
		ws.cell(row=row_num, column=8, value=row_data.get("igst"))
		ws.cell(row=row_num, column=9, value=row_data.get("total_amount"))
		
		# Format currency columns
		for col in [4, 6, 7, 8, 9]:
			ws.cell(row=row_num, column=col).number_format = '#,##0.00'
		
		# Format quantity column
		ws.cell(row=row_num, column=3).number_format = '#,##0.00'
		
		row_num += 1
	
	# Add totals row
	total_row = row_num
	ws.cell(row=total_row, column=1, value="Total")
	ws.cell(row=total_row, column=1).font = Font(bold=True)
	
	# Calculate totals
	ws.cell(row=total_row, column=3, value=f"=SUM(C2:C{total_row-1})")
	ws.cell(row=total_row, column=4, value=f"=SUM(D2:D{total_row-1})")
	ws.cell(row=total_row, column=6, value=f"=SUM(F2:F{total_row-1})")
	ws.cell(row=total_row, column=7, value=f"=SUM(G2:G{total_row-1})")
	ws.cell(row=total_row, column=8, value=f"=SUM(H2:H{total_row-1})")
	ws.cell(row=total_row, column=9, value=f"=SUM(I2:I{total_row-1})")
	
	# Format totals row
	for col in range(1, 10):
		cell = ws.cell(row=total_row, column=col)
		cell.font = Font(bold=True)
		cell.fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
		if col in [3, 4, 6, 7, 8, 9]:
			cell.number_format = '#,##0.00'
	
	# Adjust column widths
	column_widths = {
		'A': 15,  # HSN Code
		'B': 10,  # UOM
		'C': 12,  # Quantity
		'D': 18,  # Taxable Amount
		'E': 30,  # Item Tax Template
		'F': 15,  # CGST
		'G': 15,  # SGST
		'H': 15,  # IGST
		'I': 18   # Total Amount
	}
	
	for col, width in column_widths.items():
		ws.column_dimensions[col].width = width
	
	# Save file
	timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
	file_path = f"/home/ubuntu/frappe-bench/sites/{frappe.local.site}/private/files/GST_HSN_Summary_{timestamp}.xlsx"
	wb.save(file_path)
	
	print(f"\n✓ Excel file generated successfully!")
	print(f"✓ Location: {file_path}")
	print(f"✓ Total records: {len(data)}")
	
	return file_path


def get_gst_hsn_data():
	"""Fetch GST HSN Summary data from Sales Invoice Items"""
	
	query = """
		SELECT 
			sii.gst_hsn_code,
			sii.uom,
			SUM(sii.qty) as quantity,
			SUM(sii.rate * sii.qty) as taxable_amount,
			sii.item_tax_template,
			si.taxes_and_charges
		FROM 
			`tabSales Invoice Item` sii
		INNER JOIN 
			`tabSales Invoice` si ON sii.parent = si.name
		WHERE 
			si.docstatus = 1
			AND si.customer != 'Online Ecom (Swiggy/Zomato)'
			AND si.company = 'Touchstone Foundation Jaipur'
			AND si.posting_date BETWEEN '2024-04-01' AND '2025-03-31'
		GROUP BY 
			sii.gst_hsn_code, 
			sii.uom, 
			sii.item_tax_template,
			si.taxes_and_charges
		ORDER BY 
			sii.gst_hsn_code
	"""
	
	result = frappe.db.sql(query, as_dict=1)
	
	# Process the data to calculate CGST, SGST, IGST
	processed_data = []
	
	for row in result:
		taxable_amount = row.get("taxable_amount", 0)
		
		# Extract tax percentage from item_tax_template
		tax_rate = extract_tax_percentage(row.get("item_tax_template", ""))
		
		# Calculate tax amounts based on in-state or out-of-state
		cgst = 0
		sgst = 0
		igst = 0
		
		if tax_rate > 0:
			total_tax = taxable_amount * (tax_rate / 100)
			
			# Check if it's in-state or out-of-state GST
			if row.get("taxes_and_charges") == "Output In State GST":
				# Split equally between CGST and SGST
				cgst = total_tax / 2
				sgst = total_tax / 2
			else:
				# Out of state - all goes to IGST
				igst = total_tax
		
		# Calculate total amount
		total_amount = taxable_amount + cgst + sgst + igst
		
		processed_data.append({
			"gst_hsn_code": row.get("gst_hsn_code"),
			"uom": row.get("uom"),
			"quantity": row.get("quantity"),
			"taxable_amount": taxable_amount,
			"item_tax_template": row.get("item_tax_template"),
			"cgst": cgst,
			"sgst": sgst,
			"igst": igst,
			"total_amount": total_amount
		})
	
	return processed_data


def extract_tax_percentage(tax_template):
	"""
	Extract tax percentage from item tax template string.
	Example: "GST - 5% - TSFJ" returns 5.0
	"""
	if not tax_template:
		return 0
	
	# Search for pattern like "5%" or "12%" etc.
	match = re.search(r'(\d+(?:\.\d+)?)\s*%', tax_template)
	
	if match:
		return float(match.group(1))
	
	return 0

