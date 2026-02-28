import re
import frappe
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
import os
from datetime import timedelta


class ExpenseSummary:
    def __init__(self, cost_heads, from_date, to_date, recipients):
        self.cost_heads = cost_heads
        self.from_date = from_date
        self.to_date = to_date
        self.recipients = recipients
        self.path = None
        self.cost_heads_wise_summary = None
        self.account_wise_summary = None
        self.body = None
        self.load_data()

    def load_data(self):
        self.path = self.make_excel()
        if not self.path:
            return

        cost_heads_results_gl, cost_heads_total_expense_gl, cost_heads_total_income_gl = (
            self.make_cost_heads_wise_summary(is_internal_adjustment=False)
        )

        (
            cost_heads_results_internal_adjustment,
            cost_heads_total_expense_internal_adjustment,
            cost_heads_total_income_internal_adjustment,
        ) = self.make_cost_heads_wise_summary(is_internal_adjustment=True)
        
        
        

        (
            account_wise_results_gl,
            account_wise_total_expense_gl,
            account_wise_total_income_gl,
        ) = self.make_account_wise_summary(is_internal_adjustment=False)

        (
            account_wise_results_internal_adjustment,
            account_wise_total_expense_internal_adjustment,
            account_wise_total_income_internal_adjustment,
        ) = self.make_account_wise_summary(is_internal_adjustment=True)
        
        grand_total = frappe.render_template(
            "hkm/expense_summary/template/grand_total.html",
            {
                "gl_entry_total_expense": account_wise_total_expense_gl,
                "gl_entry_total_income": account_wise_total_income_gl,
                "internal_adjustment_total_expense": account_wise_total_expense_internal_adjustment,
                "internal_adjustment_total_income": account_wise_total_income_internal_adjustment,
                "grand_total_expense": account_wise_total_expense_gl + account_wise_total_expense_internal_adjustment,
                "grand_total_income": account_wise_total_income_gl + account_wise_total_income_internal_adjustment,
                
            },
        )

        cost_heads_summary_gl = frappe.render_template(
            "hkm/expense_summary/template/cost_head/cost_head.html",
            {
                "cost_heads_results": cost_heads_results_gl,
                "cost_heads_total_expense": cost_heads_total_expense_gl,
                "cost_heads_total_income": cost_heads_total_income_gl,
            },
        )
        cost_heads_summary_internal_adjustment = frappe.render_template(
            "hkm/expense_summary/template/cost_head/cost_head.html",
            {
                "cost_heads_results": cost_heads_results_internal_adjustment,
                "cost_heads_total_expense": cost_heads_total_expense_internal_adjustment,
                "cost_heads_total_income": cost_heads_total_income_internal_adjustment,
            },
        )

        account_summary_gl = frappe.render_template(
            "hkm/expense_summary/template/accounts/account.html",
            {
                "account_results": account_wise_results_gl,
                "account_total_expense": account_wise_total_expense_gl,
                "account_total_income": account_wise_total_income_gl,
            },
        )
        account_summary_internal_adjustment = frappe.render_template(
            "hkm/expense_summary/template/accounts/account.html",
            {
                "account_results": account_wise_results_internal_adjustment,
                "account_total_expense": account_wise_total_expense_internal_adjustment,
                "account_total_income": account_wise_total_income_internal_adjustment,
            },
        )

        self.body = frappe.render_template(
            "hkm/expense_summary/template/main.html",
            {
                "from_date": self.from_date.strftime("%B %d, %Y"),
                "to_date": self.to_date.strftime("%B %d, %Y"),
                "cost_heads_summary_gl": cost_heads_summary_gl,
                "cost_heads_summary_internal_adjustment": cost_heads_summary_internal_adjustment,
                "account_summary_gl": account_summary_gl,
                "account_summary_internal_adjustment": account_summary_internal_adjustment,
                "grand_total": grand_total,
            },
        )
        # print(self.body)

        frappe.sendmail(
            recipients=self.recipients,
            subject="Summary Report Expenses",
            message=self.body,
            attachments=[
                {
                    "fname": os.path.basename(self.path),
                    "fcontent": open(self.path, "rb").read(),
                }
            ],
        )
        os.remove(self.path)

    def make_account_wise_summary(self , is_internal_adjustment=False):
        gl_query = """
        SELECT 
            gl.account AS account,
            ac.root_type AS adjustment_for,
            SUM(CASE 
                WHEN ac.root_type = 'Expense' THEN gl.debit - gl.credit 
                ELSE 0 
            END) AS total_expense,
            SUM(CASE 
                WHEN ac.root_type = 'Income' THEN gl.credit - gl.debit 
                ELSE 0 
            END) AS total_income
        FROM `tabGL Entry` gl
        LEFT JOIN `tabAccount` ac ON gl.account = ac.name
        WHERE gl.cost_head IN %(cost_heads)s
            AND gl.is_cancelled = 0
            AND gl.posting_date BETWEEN %(from_date)s AND %(to_date)s
            AND ac.root_type IN ("Expense", "Income")
        GROUP BY gl.account, ac.root_type
    """

        ia_query = """
        SELECT
            iai.account AS account,
            ia.adjustment_for AS adjustment_for,
            SUM(CASE 
                WHEN ia.adjustment_for = 'Expense' THEN iai.debit - iai.credit 
                ELSE 0 
            END) AS total_expense,
            SUM(CASE 
                WHEN ia.adjustment_for = 'Income' THEN iai.credit - iai.debit 
                ELSE 0 
            END) AS total_income
        FROM `tabInternal Adjustment` ia
        JOIN `tabInternal Adjustment Detail` iai ON iai.parent = ia.name
        WHERE ia.docstatus = 1
            AND iai.cost_head IN %(cost_heads)s
            AND ia.posting_date BETWEEN %(from_date)s AND %(to_date)s
        GROUP BY iai.account, ia.adjustment_for
    """
        results = frappe.db.sql(
            gl_query if not is_internal_adjustment else ia_query,
            {
                "cost_heads": self.cost_heads,
                "from_date": self.from_date,
                "to_date": self.to_date,
            },
            as_dict=True,
        )

        total_expense = sum(row.total_expense for row in results)
        total_income = sum(row.total_income for row in results)
        results = [
            row
            for row in results
            if not (row.total_expense == 0 and row.total_income == 0)
        ]
        
        data = {}
        if results :
            for row in results:
                key = row["account"]
                if key not in data:
                    data[key] = {
                        "account": row["account"],
                        "adjustment_for": row["adjustment_for"],
                        "total_income": 0,
                        "total_expense": 0,
                    }
                data[key]["total_expense"] += row["total_expense"]
                data[key]["total_income"] += row["total_income"]
        results = list(data.values())
        return results, total_expense, total_income

    def make_cost_heads_wise_summary(self , is_internal_adjustment=False):
        gl_query = """
            SELECT 
                gl.cost_head,
                ac.root_type AS adjustment_for,
                SUM(CASE 
                    WHEN ac.root_type = 'Expense' THEN gl.debit - gl.credit 
                    ELSE 0 
                END) AS total_expense,
                SUM(CASE 
                    WHEN ac.root_type = 'Income' THEN gl.credit - gl.debit 
                    ELSE 0 
                END) AS total_income
            FROM `tabGL Entry` gl
            LEFT JOIN `tabAccount` ac ON gl.account = ac.name
            WHERE gl.cost_head IN %(cost_heads)s
                AND gl.is_cancelled = 0
                AND gl.posting_date BETWEEN %(from_date)s AND %(to_date)s
                AND ac.root_type IN ("Expense", "Income")
            GROUP BY gl.cost_head, ac.root_type
        """

    
        ia_query = """
        SELECT
            iai.cost_head AS cost_head,
            ia.adjustment_for AS adjustment_for,
            SUM(CASE 
                WHEN ia.adjustment_for = 'Expense' THEN iai.debit - iai.credit 
                ELSE 0 
            END) AS total_expense,
            SUM(CASE 
                WHEN ia.adjustment_for = 'Income' THEN iai.credit - iai.debit 
                ELSE 0 
            END) AS total_income
        FROM `tabInternal Adjustment` ia
        JOIN `tabInternal Adjustment Detail` iai ON iai.parent = ia.name
        WHERE ia.docstatus = 1
            AND iai.cost_head IN %(cost_heads)s
            AND ia.posting_date BETWEEN %(from_date)s AND %(to_date)s
        GROUP BY iai.cost_head, ia.adjustment_for
    """
        results = frappe.db.sql(
            gl_query if not is_internal_adjustment else ia_query,
            {
                "cost_heads": self.cost_heads,
                "from_date": self.from_date,
                "to_date": self.to_date,
            },
            as_dict=True,
        )
        # Calculate total expense
        total_expense = sum(row.total_expense for row in results)
        total_income = sum(row.total_income for row in results)
        results = [
                row
                for row in results
                if not (row.total_expense == 0 and row.total_income == 0)
            ]
        
        data = {}
        if results :
            for row in results:
                key = row["cost_head"]
                if key not in data:
                    data[key] = {
                        "cost_head": row["cost_head"],
                        "adjustment_for": row["adjustment_for"],
                        "total_income": 0,
                        "total_expense": 0,
                    }
                data[key]["total_expense"] += row["total_expense"]
                data[key]["total_income"] += row["total_income"]
            results = list(data.values())
        
        
        
        return results, total_expense, total_income

    def make_excel(self):
        query = """
    SELECT 
        gl.posting_date,
        gl.account,
        gl.voucher_type,
        gl.voucher_no,
        CASE 
            WHEN ac.root_type = 'Expense' THEN gl.debit - gl.credit 
            ELSE 0 
        END AS expense,
        CASE 
            WHEN ac.root_type = 'Income' THEN gl.credit - gl.debit 
            ELSE 0 
        END AS income,
        gl.cost_head,
        gl.company,
        gl.remarks
    FROM `tabGL Entry` gl
    LEFT JOIN `tabAccount` ac 
        ON gl.account = ac.name
    WHERE gl.cost_head IN %(cost_heads)s
        AND gl.is_cancelled = 0
        AND gl.posting_date BETWEEN %(from_date)s AND %(to_date)s
        AND ac.root_type IN ('Expense', 'Income')

    UNION ALL

    SELECT 
        ia.posting_date,
        iai.account,
        'Internal Adjustment' AS voucher_type,
        ia.name AS voucher_no,
        CASE 
            WHEN ia.adjustment_for = 'Expense' THEN iai.debit - iai.credit 
            ELSE 0 
        END AS expense,
        CASE 
            WHEN ia.adjustment_for = 'Income' THEN iai.credit - iai.debit 
            ELSE 0 
        END AS income,
        iai.cost_head,
        "Internal Adjustment" AS company,
        CONCAT('IAEC :- ', IFNULL(ia.remarks, ia.name)) AS remarks  # v16: TODO - Use pypika Field objects for ifnull
    FROM `tabInternal Adjustment` ia
    JOIN `tabInternal Adjustment Detail` iai
        ON iai.parent = ia.name
    WHERE ia.docstatus = 1
        AND iai.cost_head IN %(cost_heads)s
        AND ia.posting_date BETWEEN %(from_date)s AND %(to_date)s
"""
        results = frappe.db.sql(
            query,
            {
                "cost_heads": tuple(self.cost_heads),
                "from_date": self.from_date,
                "to_date": self.to_date,
            },
            as_dict=True,
        )

        total_debit = sum(row.expense for row in results if row.expense > 0) or 0
        total_credit = sum(row.income for row in results if row.income > 0) or 0

        print(total_debit, total_credit)
        if total_debit == total_credit == 0:
            return None

        def clean_text(text):
            if not text:
                return ""
            text = re.sub(r"\s+", " ", str(text))
            return text.strip()

        for gl_entry in results:
            gl_entry["remarks"] = clean_text(gl_entry.get("remarks"))
            gl_entry["url"] = frappe.utils.get_url_to_form(
                gl_entry.get("voucher_type"), gl_entry.get("voucher_no")
            )

        wb = openpyxl.Workbook()

        header_font = Font(bold=True, color="FFFFFF", size=12)
        header_fill = PatternFill(
            start_color="4F81BD", end_color="4F81BD", fill_type="solid"
        )
        cell_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )
        center_align = Alignment(horizontal="center", vertical="center")
        hyperlink_font = Font(color="0000FF", underline="single")

        fields = [
            ("Posting Date", "posting_date"),
            ("Company", "company"),
            ("Account", "account"),
            ("Voucher Type", "voucher_type"),
            ("Voucher No", "voucher_no"),
            ("Income", "income"),
            ("Expense", "expense"),
            ("Cost Head", "cost_head"),
            ("Remarks", "remarks"),
            ("URL", "url"),
        ]

        def create_sheet(sheet_title, skip_field, main_field, start_row=4):
            filtered_results = [
                entry for entry in results if entry.get(main_field) != 0
            ]

            ws = (
                wb.create_sheet(title=sheet_title)
                if sheet_title != "Expense Summary"
                else wb.active
            )
            ws.title = sheet_title

            ws.merge_cells("A1:J1")
            ws["A1"] = f"{sheet_title} ({self.from_date} to {self.to_date})"
            ws["A1"].font = Font(bold=True, size=14)
            ws["A1"].alignment = center_align
            col_idx = 1
            for field_name, field_key in fields:
                if field_name == skip_field:
                    continue
                cell = ws.cell(row=3, column=col_idx)
                cell.value = field_name
                cell.font = header_font
                cell.fill = header_fill
                cell.border = cell_border
                cell.alignment = center_align
                col_idx += 1

            for row, entry in enumerate(filtered_results, start=start_row):
                col_idx = 1
                for field_name, field_key in fields:
                    if field_name == skip_field:
                        continue
                    cell = ws.cell(row=row, column=col_idx)
                    cell.value = entry.get(field_key)
                    cell.border = cell_border

                    if field_key == "url" and entry.get("url"):
                        cell.hyperlink = entry.get("url")
                        cell.value = "View Voucher"
                        cell.font = hyperlink_font
                        cell.alignment = Alignment(horizontal="left", vertical="center")
                    elif field_key in ["posting_date", "transaction_date", "due_date"]:
                        cell.number_format = "YYYY-MM-DD"
                    elif field_key in ["income", "expense"]:
                        cell.number_format = "#,##0.00"
                    else:
                        cell.alignment = Alignment(horizontal="left", vertical="center")
                    col_idx += 1

            wider_columns = {
                "Account": 25,
                "Voucher No": 25,
                "Cost Head": 25,
                "Remarks": 30,
                "URL": 20,
            }
            for col in range(1, len(fields) - 1 + 1):
                column_letter = get_column_letter(col)
                ws.column_dimensions[column_letter].width = wider_columns.get(
                    [f[0] for f in fields if f[0] != skip_field][col - 1], 15
                )

            ws.freeze_panes = "A4"

        default_sheet = wb.active
        wb.remove(default_sheet)
        if total_debit > 0:
            create_sheet("Expense", skip_field="Income", main_field="expense")
        if total_credit > 0:
            create_sheet("Income", skip_field="Expense", main_field="income")

        file_name = f"Summary_{self.from_date}_to_{self.to_date}.xlsx"
        file_path = os.path.join(frappe.get_site_path(), "private", "files", file_name)
        wb.save(file_path)
        return file_path
